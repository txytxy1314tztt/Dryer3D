"""Extract slide-by-slide evidence from a PPTX without modifying the source deck."""

from __future__ import annotations

import argparse
import json
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


PML = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def shape_text(shape) -> str:
    chunks: list[str] = []
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                chunks.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def extract_notes(slide) -> list[str]:
    try:
        frame = slide.notes_slide.notes_text_frame
    except (AttributeError, ValueError):
        return []
    if frame is None:
        return []
    notes: list[str] = []
    for paragraph in frame.paragraphs:
        text = "".join(run.text for run in paragraph.runs).strip()
        if text and text not in {"Slide image", "Slide number"}:
            notes.append(text)
    return notes


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--json", dest="json_path", type=Path, required=True)
    parser.add_argument("--markdown", dest="md_path", type=Path, required=True)
    args = parser.parse_args()

    prs = Presentation(args.pptx)
    slides: list[dict] = []
    total_types: Counter[str] = Counter()

    for index, slide in enumerate(prs.slides, start=1):
        entries: list[dict] = []
        for shape_index, shape in enumerate(slide.shapes, start=1):
            type_name = str(shape.shape_type)
            try:
                type_name = MSO_SHAPE_TYPE(shape.shape_type).name
            except ValueError:
                pass
            total_types[type_name] += 1
            text = shape_text(shape)
            entries.append(
                {
                    "index": shape_index,
                    "name": shape.name,
                    "type": type_name,
                    "text": text,
                    "left": int(shape.left),
                    "top": int(shape.top),
                    "width": int(shape.width),
                    "height": int(shape.height),
                    "alt_title": getattr(shape, "title", "") or "",
                    "alt_description": getattr(shape, "description", "") or "",
                }
            )
        slides.append(
            {
                "number": index,
                "layout": slide.slide_layout.name,
                "notes": extract_notes(slide),
                "shapes": entries,
            }
        )

    package_counts: Counter[str] = Counter()
    package_files: list[str] = []
    with zipfile.ZipFile(args.pptx) as archive:
        package_files = archive.namelist()
        for name in package_files:
            lower = name.lower()
            if lower.startswith("ppt/media/"):
                package_counts[f"media:{Path(lower).suffix or 'none'}"] += 1
            elif lower.startswith("ppt/embeddings/"):
                package_counts["embeddings"] += 1
            elif lower.startswith("ppt/charts/") and lower.endswith(".xml"):
                package_counts["charts"] += 1
            elif lower.startswith("ppt/diagrams/"):
                package_counts["diagrams"] += 1
            elif lower.startswith("ppt/comments/"):
                package_counts["comments"] += 1

    result = {
        "source": str(args.pptx.resolve()),
        "size_bytes": args.pptx.stat().st_size,
        "slide_count": len(slides),
        "slide_size_emu": {"width": prs.slide_width, "height": prs.slide_height},
        "shape_types": dict(sorted(total_types.items())),
        "package_counts": dict(sorted(package_counts.items())),
        "package_file_count": len(package_files),
        "slides": slides,
    }
    args.json_path.parent.mkdir(parents=True, exist_ok=True)
    args.json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = [
        "# PPT 原始逐页提取",
        "",
        f"- 源文件：`{result['source']}`",
        f"- 文件大小：{result['size_bytes']} bytes",
        f"- 幻灯片数：{result['slide_count']}",
        f"- 包内对象统计：`{json.dumps(result['package_counts'], ensure_ascii=False)}`",
        "",
        "> 本文件是结构化提取证据，不等同于视觉渲染复核。对象坐标单位为 EMU。",
        "",
    ]
    for slide in slides:
        lines.extend([f"## 第 {slide['number']} 页", "", f"版式：`{slide['layout']}`", ""])
        texts = [shape for shape in slide["shapes"] if shape["text"]]
        if texts:
            for shape in texts:
                safe_text = shape["text"].replace("\n", "  \n")
                lines.extend(
                    [
                        f"### 对象 {shape['index']} · {shape['name']} · {shape['type']}",
                        "",
                        safe_text,
                        "",
                    ]
                )
        else:
            lines.extend(["（未提取到文字）", ""])
        if slide["notes"]:
            lines.extend(["### 备注", "", *slide["notes"], ""])
        object_summary = Counter(shape["type"] for shape in slide["shapes"])
        lines.extend([f"对象统计：`{dict(object_summary)}`", ""])

    args.md_path.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
